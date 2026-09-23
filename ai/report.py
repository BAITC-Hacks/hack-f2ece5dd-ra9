"""Generate the meeting protocol locally as an editable DOCX."""
from io import BytesIO


def create_docx(meeting):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    for style in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph('Протокол совещания', 'Title')
    doc.add_paragraph(meeting.get('title') or 'Совещание')
    doc.add_paragraph('Автоматический черновик. Проверьте имена, сроки и цитаты перед использованием.')
    doc.add_heading('Краткие итоги', level=1)
    for point in meeting['summary']['key_points']:
        doc.add_paragraph(point, 'List Bullet')
    if not meeting['summary']['key_points']:
        doc.add_paragraph(meeting['summary']['text'])
    doc.add_heading('Поручения', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ('Поручение', 'Ответственный', 'Срок')):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for task in meeting['tasks']:
        for cell, value in zip(table.add_row().cells,
                               (task['title'], task['responsible'] or 'Не указан', task['deadline'] or 'Не указан')):
            cell.text = value
    if not meeting['tasks']:
        doc.add_paragraph('Поручения автоматически не выделены.')
    doc.add_heading('Основания поручений', level=1)
    for task in meeting['tasks']:
        evidence = task['evidence']
        doc.add_paragraph(f"{task['id']} [{evidence['start']}–{evidence['end']}]: {evidence['quote']}")
    doc.add_heading('Транскрипт', level=1)
    names = meeting.get('speaker_names', {})
    for segment in meeting['transcript']['segments']:
        speaker = segment.get('speaker_id', segment.get('speaker', 'speaker-unknown'))
        name = names.get(speaker, 'Говорящий не определён' if speaker == 'speaker-unknown' else speaker)
        doc.add_paragraph(f"[{segment['start']:.2f}–{segment['end']:.2f}] {name}: {segment['text']}")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
