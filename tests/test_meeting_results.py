from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from ai.analyzer import analyse
from ai.diarization import assign_speakers
from ai.meeting import build_meeting, finalize_recording
from ai.report import create_docx


class MeetingTests(unittest.TestCase):
    def test_evidence_does_not_span_entire_meeting(self):
        result = analyse({'segments': [
            {'start': 0, 'end': 5, 'text': 'Обсудили бюджет. Алия, подготовь'},
            {'start': 5, 'end': 8, 'text': 'отчёт до завтра.'},
            {'start': 8, 'end': 12, 'text': 'Обсудили риски.'},
            {'start': 12, 'end': 15, 'text': 'Данияр, проверь звук до завтра.'}]})
        self.assertEqual(result['tasks'][0]['evidence']['end'], '00:08')
        self.assertEqual(result['tasks'][1]['evidence']['start'], '00:12')

    def test_keeps_two_owners_of_identical_task(self):
        result = analyse({'segments': [
            {'text': 'Алия, представь отчёт до завтра.'},
            {'text': 'Данияр, представь отчёт до завтра.'}]})
        self.assertEqual(len(result['tasks']), 2)

    def test_summary_uses_decision_instead_of_opening(self):
        result = analyse({'segments': [
            {'text': 'Коллеги, начинаем.'}, {'text': 'Прошу коротко и по существу.'},
            {'text': 'Решили перенести запуск на октябрь.'}]})
        self.assertEqual(result['summary']['key_points'], ['Решили перенести запуск на октябрь.'])

    def test_self_commitment_uses_named_speaker_and_keeps_identity(self):
        result = build_meeting([{'speaker_id': 'speaker-2', 'start': 1, 'end': 4,
                                 'text': 'Я подготовлю отчёт до завтра.'}], speaker_names={'speaker-2': 'Алия'})
        self.assertEqual(result['tasks'][0]['responsible'], 'Алия')
        self.assertEqual(result['tasks'][0]['responsible_id'], 'speaker-2')
        self.assertIn('Алия', result['summary']['text'])

    def test_unknown_speaker_is_not_an_assignee(self):
        result = build_meeting([{'speaker_id': 'speaker-unknown', 'text': 'Я подготовлю отчёт до завтра.'}])
        self.assertIsNone(result['tasks'][0]['responsible'])

    def test_self_commitment_keeps_voice_id_before_name_mapping(self):
        result = build_meeting([{'speaker_id': 'speaker-2', 'text': 'Я подготовлю отчёт до завтра.'}])
        self.assertEqual(result['tasks'][0]['responsible_id'], 'speaker-2')
        self.assertTrue(result['tasks'][0]['requires_review'])

    def test_upload_rejects_empty_unsupported_and_missing_token(self):
        from fastapi.testclient import TestClient
        from backend.app import create_app
        with TemporaryDirectory() as folder, TestClient(create_app(Path(folder))) as client:
            url = '/api/recordings/upload?filename=meeting.wav'
            self.assertEqual(client.post(url, content=b'audio').status_code, 403)
            headers = {'x-recorder-token': client.get('/api/status').json()['csrf_token']}
            self.assertEqual(client.post(url, headers=headers, content=b'').status_code, 422)
            self.assertEqual(client.post('/api/recordings/upload?filename=bad.exe',
                                         headers=headers, content=b'audio').status_code, 422)
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_upload_hands_bytes_to_local_worker_and_cleans_temporary_file(self):
        from fastapi.testclient import TestClient
        from backend.app import create_app
        class Importer:
            def ensure_idle(self):
                pass
            def import_audio(self, path, title):
                self.data, self.title = path.read_bytes(), title
                return {'id': 'f' * 32, 'status': 'processing'}
        importer = Importer()
        with TemporaryDirectory() as folder, TestClient(create_app(Path(folder), transcription=importer)) as client:
            headers = {'x-recorder-token': client.get('/api/status').json()['csrf_token']}
            response = client.post('/api/recordings/upload?filename=meeting.wav&title=Demo',
                                   headers=headers, content=b'local audio bytes')
            self.assertEqual(response.status_code, 202)
            self.assertEqual(importer.data, b'local audio bytes')
            self.assertEqual(importer.title, 'Demo')
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_diarization_keeps_ambiguous_segments_unknown(self):
        turns = [{'start': 0, 'end': 2, 'speaker_id': 'A'}, {'start': 2, 'end': 4, 'speaker_id': 'B'}]
        result = assign_speakers([{'start': 0, 'end': 1, 'text': 'А'},
                                  {'start': 1, 'end': 3, 'text': 'Два голоса'},
                                  {'start': 3, 'end': 4, 'text': 'Б'}], turns)
        self.assertEqual([s['speaker_id'] for s in result], ['speaker-1', 'speaker-unknown', 'speaker-2'])

    @patch('ai.diarization.diarize_recording', side_effect=RuntimeError('model unavailable'))
    def test_missing_model_preserves_transcript(self, _model):
        result = finalize_recording('missing.wav', [{'text': 'Текст', 'start': 1, 'end': 2}], model_path='missing')
        self.assertEqual(result['diarization_status'], 'error')
        self.assertEqual(result['transcript']['segments'][0]['text'], 'Текст')
        self.assertTrue(result['warnings'])

    def test_docx_preserves_kazakh_and_task_fields(self):
        from docx import Document
        result = build_meeting([{'speaker_id': 'speaker-1', 'start': 1, 'end': 5,
                                 'text': 'Алия, есепті дайындаңыз, мерзімі ертеңге дейін.'}])
        document = Document(BytesIO(create_docx(result)))
        self.assertEqual(document.tables[0].rows[1].cells[1].text, 'Алия')
        self.assertEqual(document.tables[0].rows[1].cells[2].text, 'ертеңге дейін')
        self.assertTrue(any('есепті дайындаңыз' in p.text for p in document.paragraphs))

    def test_completed_result_api_docx_and_names(self):
        from fastapi.testclient import TestClient
        from backend.app import create_app
        from backend.recorder import save_json
        sid = 'e' * 32
        with TemporaryDirectory() as folder:
            directory = Path(folder) / sid
            directory.mkdir()
            save_json(directory / 'session.json', {'id': sid, 'status': 'completed'})
            meeting = build_meeting([{'speaker_id': 'speaker-1', 'start': 0, 'end': 3,
                                      'text': 'Я подготовлю отчёт до завтра.'}])
            with TestClient(create_app(Path(folder))) as client:
                base = f'/api/sessions/{sid}'
                self.assertEqual(client.get(base + '/result').status_code, 409)
                save_json(directory / 'result.json', meeting)
                self.assertEqual(client.get(base + '/result').status_code, 200)
                token = client.get('/api/status').json()['csrf_token']
                self.assertEqual(client.post(base + '/speakers', json={'names': {}}).status_code, 403)
                response = client.post(base + '/speakers', headers={'x-recorder-token': token},
                                       json={'names': {'speaker-1': 'Алия'}})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()['tasks'][0]['responsible'], 'Алия')
                self.assertTrue(client.get(base + '/protocol.docx').content.startswith(b'PK'))


if __name__ == '__main__':
    unittest.main()
